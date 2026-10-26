from datetime import date
from decimal import Decimal
import json
import io
import zipfile
import tempfile
import time
from pathlib import Path
from unittest.mock import patch
from django.contrib.auth.models import Group, User
from django.core.cache import cache
from django.core.management import call_command, CommandError
from django.test import Client, TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from . import models as m
from .legacy_ui import PAGES
from . import db_admin
from .forms import REGISTRY, UserForm, form_for
from .estd import route_card_data, render_route_card


class WorkflowTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("import_sample_data", verbosity=0)
        # Тестовые пароли задаются только во временной БД тестового прогона.
        for username in ("admin", "user", "planner", "manager", "supervisor"):
            account = User.objects.get(username=username)
            account.is_active = True
            account.set_password(username)
            if username == "admin":
                account.is_staff = True
                account.is_superuser = True
            account.save()

    def setUp(self):
        self.client = Client(HTTP_HOST="127.0.0.1")
        self.client.login(username="admin", password="admin")

    def test_calculations_and_demo_data(self):
        route = m.RouteOperation.objects.get(product__name="Патрон станка", number=10)
        self.assertEqual(route.batch_hours(12), Decimal("28.500"))
        self.assertEqual(route.calculated_piece_hours(12), Decimal("2.375"))
        order = m.Order.objects.get(number="З-2026-041")
        self.assertEqual(order.progress, 42)
        requirement = m.MaterialRequirement.objects.get(order_item__order=order)
        self.assertEqual(requirement.required, Decimal("99.000"))
        self.assertEqual(requirement.deficit, 0)
        line = m.ProgramLine.objects.get(order_item__order=order)
        self.assertEqual(line.forecast_date, date(2026, 9, 30))

    def test_order_form_persists_and_audits(self):
        response = self.client.post("/records/orders/new/", {
            "number": "З-TEST-1", "customer": m.Customer.objects.first().pk,
            "department": m.Department.objects.first().pk, "request_number": "T-1",
            "start_date": "2026-10-15", "due_date": "2026-11-15", "priority": 2,
            "status": "draft", "constraint": "", "note": "Тест",
        })
        self.assertEqual(response.status_code, 302)
        order = m.Order.objects.get(number="З-TEST-1")
        self.assertEqual(order.note, "Тест")
        self.assertTrue(m.AuditLog.objects.filter(model_name="Order", object_id=order.pk, action="create").exists())

    def test_job_generation_and_program_revision(self):
        line = m.ProgramLine.objects.first()
        before = m.Job.objects.count()
        response = self.client.post(f"/program-lines/{line.pk}/generate-jobs/")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(m.Job.objects.count(), before)
        program = line.program
        response = self.client.post(f"/programs/{program.pk}/revise/", {"reason": "Изменение сроков"})
        self.assertEqual(response.status_code, 302)
        revision = m.Program.objects.get(name=program.name, version=2)
        self.assertEqual(revision.status, "draft")
        self.assertEqual(revision.lines.count(), program.lines.count())

    def test_delivery_requires_acceptance(self):
        job = m.Job.objects.get(program_line__order_item__order__number="З-2026-041")
        order = job.program_line.order_item.order
        response = self.client.post("/records/events/new/", {
            "order": order.pk, "job": job.pk, "kind": "delivered",
            "quantity": 6, "date": "2026-10-14", "document": "Тест",
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Нельзя сдать больше принятого количества")
        self.assertFalse(m.ProductionEvent.objects.filter(document="Тест").exists())

    def test_acceptance_cannot_exceed_quality_approval(self):
        job = m.Job.objects.get(program_line__order_item__order__number="З-2026-041")
        response = self.client.post("/records/events/new/", {
            "order": job.program_line.order_item.order_id, "job": job.pk,
            "kind": "accepted", "quantity": 1, "date": "2026-10-14", "document": "ОТК-дубль",
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Принятое количество должно пройти контроль качества")

    def test_role_blocks_changes(self):
        self.client.logout()
        self.client.login(username="user", password="user")
        self.assertEqual(self.client.get("/records/orders/").status_code, 200)
        self.assertEqual(self.client.get("/records/jobs/new/").status_code, 403)

    def test_pdf_and_xlsx_exports(self):
        self.assertEqual(self.client.get("/reports/export.pdf")["Content-Type"], "application/pdf")
        self.assertEqual(self.client.get("/reports/export/")["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_material_file_import_updates_postgresql(self):
        upload = SimpleUploadedFile("materials.csv", "code;name;unit;stock_quantity;reserved_quantity\nM-TEST;Лист тестовый;кг;15,5;3,0\n".encode("utf-8"), content_type="text/csv")
        response = self.client.post("/materials/import/", {"file": upload})
        self.assertEqual(response.status_code, 302)
        material = m.Material.objects.get(code="M-TEST")
        self.assertEqual(material.available_quantity, Decimal("12.500"))
        self.assertTrue(m.AuditLog.objects.filter(model_name="Material", object_id=material.pk, action="import").exists())
        invalid = SimpleUploadedFile("bad.xlsx", b"not a workbook")
        error = self.client.post("/materials/import/", {"file": invalid}, HTTP_ACCEPT="application/json")
        self.assertEqual(error.status_code, 400)
        self.assertFalse(error.json()["ok"])

    def test_material_reservation_updates_both_totals(self):
        requirement = m.MaterialRequirement.objects.get(order_item__order__number="З-2026-042")
        old_material_reserved = requirement.material.reserved_quantity
        response = self.client.post(f"/requirements/{requirement.pk}/reserve/", {"quantity": "10", "reason": "Для партии"})
        self.assertEqual(response.status_code, 302)
        requirement.refresh_from_db()
        requirement.material.refresh_from_db()
        self.assertEqual(requirement.reserved, Decimal("110.000"))
        self.assertEqual(requirement.material.reserved_quantity, old_material_reserved + 10)
        response = self.client.post(f"/requirements/{requirement.pk}/reserve/", {"quantity": "1000"})
        requirement.refresh_from_db()
        self.assertEqual(requirement.reserved, Decimal("110.000"))

    def test_original_interface_and_cnc_order_save(self):
        page = self.client.get("/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "/static/production/legacy/css/style.css")
        self.assertContains(page, "/static/production/legacy/js/app.js")
        self.assertContains(page, "/static/production/bridge/server-bridge.js")
        self.assertEqual(self.client.get("/static/production/legacy/js/app.js").status_code, 200)
        response = self.client.post("/api/legacy/orders/", data=json.dumps({
            "number": "TEST-CNC-1", "customer": "ОРМ", "start_date": "2026-10-15",
            "due_date": "2026-11-15", "priority": 1, "machine": "Тестовый станок",
            "cnc": "yes", "cnc_control": "Fanuc", "cnc_program": "O1201",
            "one_c_code": "1C-100", "items": [{"name": "Тестовая деталь", "quantity": 3, "drawing": "D-1"}],
        }), content_type="application/json")
        self.assertEqual(response.status_code, 200, response.content)
        item = m.OrderItem.objects.get(order__number="TEST-CNC-1")
        self.assertEqual(item.equipment.name, "Тестовый станок")
        self.assertTrue(item.equipment.cnc)
        self.assertEqual(item.equipment.cnc_control, "Fanuc")
        self.assertEqual(item.cnc_program, "O1201")
        self.assertEqual(item.one_c_code, "1C-100")
        self.assertContains(self.client.get("/api/legacy/orders/"), "TEST-CNC-1")

    def test_order_edit_and_item_delete_preserve_other_items(self):
        created = self.client.post("/api/legacy/orders/", data=json.dumps({
            "number": "TEST-MULTI", "customer": "ОРМ", "start_date": "2026-10-15",
            "due_date": "2026-11-15", "items": [
                {"name": "Деталь А", "quantity": 2}, {"name": "Деталь Б", "quantity": 3}],
        }), content_type="application/json")
        self.assertEqual(created.status_code, 200, created.content)
        order = m.Order.objects.get(number="TEST-MULTI")
        item = order.items.get(product__name="Деталь А")
        changed = self.client.put(f"/api/legacy/orders/{order.pk}/", data=json.dumps({
            "item_id": item.pk, "number": "TEST-MULTI", "customer": "РИК",
            "start_date": "2026-10-16", "due_date": "2026-11-20",
            "product": "Деталь А", "quantity": 4, "reason": "Исправление заказа",
        }), content_type="application/json")
        self.assertEqual(changed.status_code, 200, changed.content)
        order.refresh_from_db()
        self.assertEqual(order.customer.name, "РИК")
        self.assertEqual(order.start_date, date(2026, 10, 16))
        self.assertTrue(order.items.filter(pk=item.pk).exists())
        removed = self.client.delete(f"/api/legacy/orders/{order.pk}/", data=json.dumps({"item_id": item.pk, "reason": "Ошибка ввода"}), content_type="application/json")
        self.assertEqual(removed.status_code, 200, removed.content)
        order.refresh_from_db()
        self.assertFalse(order.archived)
        self.assertEqual(order.items.count(), 1)

    def test_legacy_table_saves_in_postgresql(self):
        rows = [["020", "Токарная", "Токарные", "Тестовый станок", "4", "нормо-час", "020", "Да"]]
        response = self.client.put("/api/legacy/table/operations/", data=json.dumps({"rows": rows}), content_type="application/json")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(m.LegacyTable.objects.get(key="operations").rows, rows)
        self.assertEqual(self.client.get("/api/legacy/table/operations/").json()["rows"], rows)
        self.assertGreaterEqual(len(self.client.get("/api/legacy/table/operations/?include_inactive=1").json()["rows"]), 1)
        settings = self.client.put("/api/legacy/table/settings-groups/", data=json.dumps({"rows": [["Группа 1"]]}), content_type="application/json")
        self.assertEqual(settings.status_code, 200)
        self.assertEqual(m.LegacyTable.objects.get(key="settings-groups").rows, [["Группа 1"]])

    def test_editable_tabs_add_edit_and_remove_rows(self):
        samples = {
            "operations": [["090", "Контрольная", "Контроль", "Стол", "4", "нормо-час", "090", "Да"]],
            "personnel": [["099", "Тестовый работник", "Механический цех", "Контролер", "4", "1", "", "Работает"]],
            "settings-groups": [["Новая группа"]],
            "settings-staff": [["Новый исполнитель"]],
            "settings-periods": [["2027", "10", "20", "30", "40"]],
            "settings-columns": [["0", "Название изделия", False]],
            "settings-plan-inputs": [["0", "3"]],
        }
        for key, rows in samples.items():
            endpoint = f"/api/legacy/table/{key}/"
            with self.subTest(tab=key):
                created = self.client.put(endpoint, data=json.dumps({"rows": rows}), content_type="application/json")
                self.assertEqual(created.status_code, 200, created.content)
                self.assertEqual(self.client.get(endpoint).json()["rows"], rows)
                edited = [row.copy() for row in rows]
                edited[0][1 if len(edited[0]) > 1 else 0] = "Исправлено"
                changed = self.client.put(endpoint, data=json.dumps({"rows": edited}), content_type="application/json")
                self.assertEqual(changed.status_code, 200, changed.content)
                self.assertEqual(self.client.get(endpoint).json()["rows"], edited)
                cleared = self.client.put(endpoint, data=json.dumps({"rows": []}), content_type="application/json")
                self.assertEqual(cleared.status_code, 200, cleared.content)
                self.assertEqual(self.client.get(endpoint).json()["rows"], [])

    def test_reference_tabs_update_linked_models(self):
        operations = [["010", "Токарная", "Механообработка", "", "5", "нормо-час", "010", "Да"]]
        response = self.client.put("/api/legacy/table/operations/",
            data=json.dumps({"rows": operations}), content_type="application/json")
        self.assertEqual(response.status_code, 200, response.content)
        definition = m.OperationDefinition.objects.get(code="010")
        route = m.RouteOperation.objects.get(product__name="Патрон станка", number=10)
        self.assertEqual(route.definition, definition)

        personnel = [["001", "Иванов обновлен", "Механический цех", "Токарь", "5", "2", "21.10.2026", "Работает"]]
        response = self.client.put("/api/legacy/table/personnel/",
            data=json.dumps({"rows": personnel}), content_type="application/json")
        self.assertEqual(response.status_code, 200, response.content)
        employee = m.Employee.objects.get(personnel_number="001")
        self.assertEqual(employee.name, "Иванов обновлен")
        self.assertEqual(employee.shift, 2)
        self.assertEqual(self.client.get("/api/legacy/table/personnel/").json()["rows"], personnel)

        timesheet = [["001", "Иванов обновлен", "Токарь", "5", "160", "150", "10", "0", "165", "1,10", "93,8%"]]
        response = self.client.put("/api/legacy/table/timesheet/?month=2026-10",
            data=json.dumps({"rows": timesheet}), content_type="application/json")
        self.assertEqual(response.status_code, 200, response.content)
        entry = m.TimesheetEntry.objects.get(employee=employee, month=date(2026, 10, 1))
        self.assertEqual(entry.actual_hours, Decimal("150"))
        self.assertEqual(self.client.get("/api/legacy/table/timesheet/?month=2026-10").json()["rows"], timesheet)
        self.assertEqual(self.client.get("/api/legacy/table/timesheet/?month=2026-10").json()
            ["departments"]["001"], "Механический цех")
        invalid = self.client.put("/api/legacy/table/timesheet/?month=2026-10",
            data=json.dumps({"rows": [["999", "Неизвестный", "", "", "160", "10", "0", "0", "0"]]}),
            content_type="application/json")
        self.assertEqual(invalid.status_code, 400)
        self.assertEqual(m.TimesheetEntry.objects.filter(month=date(2026, 10, 1)).count(), 1)

    def test_entry_product_list_fields_persist_and_validate(self):
        payload = {"number": "TEST-LIST", "customer": "ОРМ", "start_date": "2026-10-21", "due_date": "2026-11-21",
            "registration_date": "2026-10-19", "request_date": "2026-10-18",
            "items": [{"name": "Изделие А", "quantity": 2, "drawing": "A-1", "drawing_count": 3, "object_name": "Корпус"},
                      {"name": "Изделие Б", "quantity": 4, "drawing_count": 1, "object_name": "Узел"}]}
        response = self.client.post("/api/legacy/orders/", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(m.Order.objects.get(number="TEST-LIST").registration_date, date(2026, 10, 19))
        self.assertEqual(m.Order.objects.get(number="TEST-LIST").request_date, date(2026, 10, 18))
        items = m.Order.objects.get(number="TEST-LIST").items.order_by("pk")
        self.assertEqual([(item.quantity, item.drawing_count, item.object_name) for item in items],
            [(2, 3, "Корпус"), (4, 1, "Узел")])
        self.assertEqual(len([item for item in self.client.get("/api/legacy/orders/").json()["orders"]
            if item["number"] == "TEST-LIST"]), 2)
        first = items.first()
        response = self.client.put(f"/api/legacy/orders/{first.order_id}/", data=json.dumps({
            "item_id": first.pk, "drawing_count": 5, "object_name": "Обновленный корпус",
            "registration_date": "2026-10-20", "request_date": "2026-10-17",
            "reason": "Исправление списка изделий"}), content_type="application/json")
        self.assertEqual(response.status_code, 200, response.content)
        first.refresh_from_db()
        self.assertEqual((first.drawing_count, first.object_name), (5, "Обновленный корпус"))
        first.order.refresh_from_db()
        self.assertEqual(first.order.registration_date, date(2026, 10, 20))
        self.assertEqual(first.order.request_date, date(2026, 10, 17))
        payload["number"] = "TEST-BAD-LIST"
        payload["items"][1]["drawing_count"] = 0
        response = self.client.post("/api/legacy/orders/", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(m.Order.objects.filter(number="TEST-BAD-LIST").exists())

    def test_automatic_assignment_uses_department_grade_and_capacity(self):
        job = m.Job.objects.get(employee__isnull=True)
        response = self.client.post(f"/jobs/{job.pk}/assign/")
        self.assertEqual(response.status_code, 302)
        job.refresh_from_db()
        self.assertEqual(job.employee.department, job.operation.work_center.department)
        self.assertTrue(m.AuditLog.objects.filter(model_name="Job", object_id=job.pk, action="assign").exists())
        job.employee = None
        job.save(update_fields=["employee"])
        definition = m.OperationDefinition.objects.create(code="ASSY", name="Сборка", kind="Сборка", min_grade=6)
        job.operation.definition = definition
        job.operation.save(update_fields=["definition"])
        response = self.client.post(f"/jobs/{job.pk}/assign/")
        self.assertEqual(response.status_code, 302)
        job.refresh_from_db()
        self.assertIsNone(job.employee)

    def test_approved_cutting_norm_changes_planned_hours(self):
        route = m.RouteOperation.objects.get(product__name="Патрон станка", number=10)
        previous = route.batch_hours(12)
        norm = m.CuttingNorm(operation=route, material=m.Material.objects.first(), tool="Резец",
            source_document="Тестовый локальный норматив", cutting_speed_m_min=Decimal("100"),
            diameter_mm=Decimal("40"), feed_mm_rev=Decimal("0.2"), teeth=1,
            pass_length_mm=Decimal("100"), passes=2, approved=False)
        norm.full_clean()
        norm.save()
        self.assertEqual(route.batch_hours(12), previous)
        self.assertGreater(norm.spindle_rpm, 790)
        self.assertGreater(norm.feed_mm_min, 150)
        norm.approved = True
        norm.full_clean()
        norm.save()
        self.assertEqual(norm.machine_hours, Decimal("0.021"))
        self.assertEqual(route.batch_hours(12), Decimal("1.752"))

    def test_draft_estd_documents_use_linked_data(self):
        product = m.Product.objects.get(name="Патрон станка")
        item = m.OrderItem.objects.get(product=product)
        for url in (f"/products/{product.pk}/operation-cards.pdf",
                    f"/products/{product.pk}/tooling-sheet.pdf",
                    f"/items/{item.pk}/kit-sheet.pdf"):
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response["Content-Type"], "application/pdf")

    def test_order_edit_does_not_change_shared_product(self):
        product = m.Product.objects.first()
        first_item = m.OrderItem.objects.filter(product=product).first()
        other_order = m.Order.objects.exclude(pk=first_item.order_id).first()
        second_item = m.OrderItem.objects.create(order=other_order, product=product, quantity=1)
        original_name = product.name
        response = self.client.put(f"/api/legacy/orders/{first_item.order_id}/", data=json.dumps({
            "item_id": first_item.pk, "product": "Новое изделие для заказа",
            "drawing": "НОВЫЙ-001", "reason": "Изменение позиции",
        }), content_type="application/json")
        self.assertEqual(response.status_code, 200, response.content)
        first_item.refresh_from_db()
        second_item.refresh_from_db()
        product.refresh_from_db()
        self.assertEqual(product.name, original_name)
        self.assertEqual(second_item.product_id, product.pk)
        self.assertNotEqual(first_item.product_id, product.pk)

    def test_duplicate_order_number_returns_validation_error(self):
        original = m.Order.objects.first()
        before = m.Order.objects.count()
        response = self.client.post("/api/legacy/orders/", data=json.dumps({
            "number": original.number, "customer": "Повторный заказчик",
            "start_date": "2026-10-15", "due_date": "2026-11-15",
            "items": [{"name": "Дубликат", "quantity": 1}],
        }), content_type="application/json")
        self.assertEqual(response.status_code, 400, response.content)
        self.assertEqual(m.Order.objects.count(), before)
        self.assertFalse(m.Customer.objects.filter(name="Повторный заказчик").exists())

    def test_order_edit_cannot_reduce_approved_plan_quantity(self):
        line = m.ProgramLine.objects.select_related("order_item__order").first()
        item = line.order_item
        response = self.client.put(f"/api/legacy/orders/{item.order_id}/", data=json.dumps({
            "item_id": item.pk, "quantity": line.quantity - 1,
            "reason": "Проверка ограничения",
        }), content_type="application/json")
        self.assertEqual(response.status_code, 400, response.content)
        item.refresh_from_db()
        self.assertGreaterEqual(item.quantity, line.quantity)

    def test_all_tpp_pages_and_login_are_served(self):
        self.client.logout()
        login = self.client.get("/login.html")
        self.assertEqual(login.status_code, 200)
        self.assertContains(login, "/static/production/legacy/css/style.css")
        self.assertContains(login, 'name="csrfmiddlewaretoken"')
        self.client.login(username="admin", password="admin")
        for name in PAGES:
            response = self.client.get(f"/{name}.html")
            self.assertEqual(response.status_code, 200, name)
            self.assertContains(response, "/static/production/legacy/js/app.js")

    def test_tabs_enforce_server_side_roles(self):
        self.client.logout()
        self.client.login(username="user", password="user")
        self.assertEqual(self.client.get("/admin.html").status_code, 403)
        self.assertEqual(self.client.get("/plan-b.html").status_code, 403)
        self.assertEqual(self.client.get("/entry.html").status_code, 200)
        self.client.logout()
        self.client.login(username="planner", password="planner")
        plan_page = self.client.get("/plan-b.html")
        self.assertEqual(plan_page.status_code, 200)
        self.assertIn('"role": "planner"', plan_page.content.decode())
        self.assertEqual(self.client.get("/personnel.html").status_code, 200)
        self.assertEqual(self.client.get("/api/legacy/table/personnel/").status_code, 200)
        self.assertEqual(self.client.put("/api/legacy/table/personnel/",
            data=json.dumps({"rows": []}), content_type="application/json").status_code, 403)

    def test_settings_permissions_match_visible_tabs(self):
        self.client.logout()
        self.client.login(username="user", password="user")
        for key in ("settings-columns", "settings-groups"):
            response = self.client.put(f"/api/legacy/table/{key}/",
                data=json.dumps({"rows": [["0", "Тест", True]]}), content_type="application/json")
            self.assertEqual(response.status_code, 200, response.content)
        for key in ("settings-staff", "settings-periods", "settings-plan-inputs"):
            response = self.client.put(f"/api/legacy/table/{key}/",
                data=json.dumps({"rows": [["0", "Тест"]]}), content_type="application/json")
            self.assertEqual(response.status_code, 403, response.content)
            self.assertEqual(self.client.get(f"/api/legacy/table/{key}/").status_code, 403)
        self.assertEqual(self.client.get("/api/legacy/table/personnel/").status_code, 403)

    def test_all_working_tabs_load_their_database_data(self):
        endpoints = [
            ("/api/legacy/dashboard/", None),
            ("/api/legacy/orders/", None),
            ("/api/legacy/materials/", None),
            ("/api/legacy/processes/", None),
            ("/api/legacy/deviations/", None),
            ("/api/legacy/report/?kind=plan-a", None),
            ("/api/legacy/report/?kind=plan-b", None),
            ("/api/legacy/report/?kind=standards", None),
            ("/api/legacy/report/?kind=balance", None),
            ("/api/legacy/report/?kind=production", None),
            ("/api/legacy/plans/?kind=annual&year=2026", None),
            ("/api/legacy/plans/?kind=a&year=2026&quarter=4", None),
            ("/api/legacy/plans/?kind=b&year=2026&quarter=4", None),
            ("/api/legacy/norming/?year=2026&quarter=4", None),
            ("/api/legacy/table/operations/", None),
            ("/api/legacy/table/personnel/", None),
            ("/api/legacy/table/timesheet/", None),
            ("/api/legacy/table/settings-groups/", None),
            ("/api/legacy/table/settings-staff/", None),
            ("/api/legacy/table/settings-periods/", None),
            ("/api/legacy/table/settings-columns/", None),
            ("/api/legacy/table/settings-plan-inputs/", None),
            ("/api/admin/database/", None),
            ("/api/admin/backup/", None),
            ("/api/admin/users/", None),
        ]
        for endpoint, _ in endpoints:
            with self.subTest(endpoint=endpoint):
                response = self.client.get(endpoint)
                self.assertEqual(response.status_code, 200, response.content[:500])
                self.assertIsInstance(response.json(), dict)

    def test_all_orm_sections_and_forms_render(self):
        for key, (model, _, _) in REGISTRY.items():
            with self.subTest(section=key):
                listing = self.client.get(f"/records/{key}/")
                self.assertEqual(listing.status_code, 200)
                form = self.client.get(f"/records/{key}/new/")
                self.assertEqual(form.status_code, 200)
                existing = model.objects.first()
                if existing and key not in {"lines", "requirements"}:
                    edit = self.client.get(f"/records/{key}/{existing.pk}/edit/")
                    self.assertEqual(edit.status_code, 200)

    def test_existing_orm_forms_save_every_editable_field(self):
        for key, (model, _, _) in REGISTRY.items():
            if key in {"lines", "requirements"}:
                continue  # Approved lines and reserved requirements need their dedicated workflows.
            obj = model.objects.first()
            if not obj:
                continue
            with self.subTest(section=key):
                form = form_for(key, instance=obj)
                data = {"change_reason": "Проверка полей формы"}
                for name, bound in [(name, form[name]) for name in form.fields]:
                    value = bound.value()
                    if value is None or value == "":
                        continue
                    if getattr(form.fields[name].widget, "input_type", "") == "checkbox":
                        if value:
                            data[name] = "on"
                    else:
                        data[name] = value
                response = self.client.post(f"/records/{key}/{obj.pk}/edit/", data)
                self.assertEqual(response.status_code, 302, response.content[:500])
                self.assertTrue(m.AuditLog.objects.filter(model_name=model.__name__, object_id=obj.pk,
                    action="update", reason="Проверка полей формы").exists())

    def test_order_to_plan_job_acceptance_and_report_workflow(self):
        created = self.client.post("/api/legacy/orders/", data=json.dumps({
            "number": "З-RELEASE-FLOW", "customer": "Тестовый заказчик",
            "start_date": "2026-10-20", "due_date": "2026-12-15", "priority": 2,
            "items": [{"name": "Тестовое изделие релиза", "drawing": "REL-001", "quantity": 2}],
        }), content_type="application/json")
        self.assertEqual(created.status_code, 200, created.content)
        order = m.Order.objects.get(number="З-RELEASE-FLOW")
        item = order.items.get()
        center = m.WorkCenter.objects.first()
        route = self.client.post("/api/legacy/processes/", data=json.dumps({
            "product_id": item.product_id, "work_center_id": center.pk,
            "number": 10, "name": "Финишная операция", "kind": "machining",
            "setup_hours": "1.00", "piece_hours": "0.500",
            "quality_required": False,
        }), content_type="application/json")
        self.assertEqual(route.status_code, 200, route.content)
        route_pk = route.json()["id"]
        self.assertEqual(self.client.put(f"/api/legacy/processes/{route_pk}/").status_code, 200)
        planned = self.client.post("/api/legacy/plans/", data=json.dumps({
            "kind": "b", "year": 2026, "quarter": 4,
        }), content_type="application/json")
        self.assertEqual(planned.status_code, 200, planned.content)
        line = m.ProgramLine.objects.get(program__name="План Б 2026/4", order_item=item)
        self.assertEqual(line.planned_hours, Decimal("2.000"))
        route_update = self.client.put(f"/api/legacy/processes/{route_pk}/", data=json.dumps({
            "product_id": item.product_id, "work_center_id": center.pk,
            "number": 10, "name": "Финишная операция", "kind": "machining",
            "setup_hours": "1.00", "piece_hours": "0.750",
            "quality_required": False,
        }), content_type="application/json")
        self.assertEqual(route_update.status_code, 200, route_update.content)
        self.assertEqual(line.planned_hours, Decimal("2.500"))
        norm_row = next(row for row in self.client.get("/api/legacy/orders/").json()["orders"]
            if row["number"] == order.number)
        self.assertEqual(norm_row["norm_per_item"], "1.250")
        self.assertEqual(norm_row["planned_hours"], "2.500")
        self.assertEqual(self.client.post(f"/program-lines/{line.pk}/generate-jobs/").status_code, 302)
        job = m.Job.objects.get(program_line=line, operation_id=route_pk)
        job.employee = m.Employee.objects.first()
        job.actual_quantity = 2
        job.actual_hours = Decimal("2.25")
        job.actual_end = date(2026, 12, 10)
        job.status = "done"
        job.full_clean()
        job.save()
        accepted = self.client.post("/records/events/new/", {
            "order": order.pk, "job": job.pk, "kind": "accepted", "quantity": 2,
            "date": "2026-12-10", "document": "ОТК-RELEASE",
        })
        self.assertEqual(accepted.status_code, 302, accepted.content)
        order.refresh_from_db()
        self.assertEqual(order.progress, 100)
        norm_row = next(row for row in self.client.get("/api/legacy/orders/").json()["orders"]
            if row["number"] == order.number)
        self.assertEqual(norm_row["actual_hours"], "2.25")
        self.assertIn(job.employee.personnel_number, norm_row["executors"])
        production = self.client.get("/api/legacy/report/?kind=production").json()
        row = next(row for row in production["rows"] if row[0] == order.number)
        self.assertEqual(row[4:7], [2, 0, 100])
        plan_b = self.client.get("/api/legacy/report/?kind=plan-b").json()
        line_row = next(row for row in plan_b["rows"] if row[0] == order.number)
        self.assertEqual(line_row[4:6], [2.5, 2.25])

    def test_norming_preview_approval_and_validation(self):
        order = m.Order.objects.filter(archived=False).first()
        period = {"year": order.due_date.year, "quarter": (order.due_date.month - 1) // 3 + 1}
        preview = self.client.get("/api/legacy/norming/", period)
        self.assertEqual(preview.status_code, 200, preview.content)
        self.assertGreater(preview.json()["items"], 0)
        self.assertIsNone(preview.json()["approval"])
        payload = {**period, "hours_per_employee": "128", "labor_factor": "1.25", "staff_count": 4}
        saved = self.client.post("/api/legacy/norming/", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(saved.status_code, 200, saved.content)
        approval = m.NormingApproval.objects.get(**period)
        self.assertEqual(approval.planned_hours, (Decimal(saved.json()["base_hours"]) * Decimal("1.25")).quantize(Decimal("0.01")))
        self.assertEqual(approval.capacity_hours, Decimal("512.00"))
        self.assertEqual(self.client.get("/api/legacy/norming/", period).json()["approval"]["approved_by"], "admin")
        self.assertEqual(self.client.post("/api/legacy/norming/", data=json.dumps({**payload,
            "staff_count": 0}), content_type="application/json").status_code, 400)
        self.assertEqual(m.NormingApproval.objects.get(**period).staff_count, 4)
        self.assertEqual(self.client.post("/api/legacy/norming/", data=json.dumps({**payload,
            "year": 2099}), content_type="application/json").status_code, 400)

    def test_norming_role_permissions(self):
        order = m.Order.objects.filter(archived=False).first()
        period = {"year": order.due_date.year, "quarter": (order.due_date.month - 1) // 3 + 1}
        user = User.objects.create_user("economist_test", password="Unique testing password 2026")
        user.groups.add(Group.objects.get(name="economist"))
        self.client.force_login(user)
        self.assertEqual(self.client.get("/api/legacy/norming/", period).status_code, 200)
        self.assertEqual(self.client.post("/api/legacy/norming/", data=json.dumps({**period,
            "hours_per_employee": "128", "labor_factor": "1", "staff_count": 4}),
            content_type="application/json").status_code, 403)

    def test_report_filters_match_order_product_department_and_dates(self):
        order = m.Order.objects.filter(archived=False).first()
        item = order.items.first()
        filters = {"kind": "production", "order": order.number,
            "product": item.product.name, "department": order.department.name,
            "from": order.due_date.isoformat(), "to": order.due_date.isoformat()}
        response = self.client.get("/api/legacy/report/", filters)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(response.json()["rows"])
        self.assertTrue(all(row[0] == order.number and row[2] == item.product.name
            for row in response.json()["rows"]))
        no_matches = self.client.get("/api/legacy/report/", {**filters, "order": "NOT-AN-ORDER"})
        self.assertEqual(no_matches.json()["rows"], [])
        self.assertEqual(self.client.get("/api/legacy/report/", {**filters,
            "from": "2026-12-31", "to": "2026-01-01"}).status_code, 400)
        export = self.client.get("/api/legacy/report/export/", filters)
        self.assertEqual(export.status_code, 200, export.content[:100])
        workbook = load_workbook(io.BytesIO(export.content), read_only=True)
        self.assertEqual(workbook.active.max_row, len(response.json()["rows"]) + 1)

    def test_admin_database_check_queries_postgresql(self):
        response = self.client.get("/api/admin/database/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["engine"], "PostgreSQL")
        self.assertTrue(response.json()["ok"])
        self.assertTrue(response.json()["database"].startswith("test_"))
        self.client.logout()
        self.client.login(username="user", password="user")
        self.assertEqual(self.client.get("/api/admin/database/").status_code, 403)

    def test_admin_users_are_saved_in_database(self):
        payload = {"username": "release_user", "fullName": "Тестовый Сотрудник", "role": "technologist", "password": "securepass123"}
        created = self.client.post("/api/admin/users/", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(created.status_code, 200, created.content)
        user = User.objects.get(username="release_user")
        self.assertTrue(user.check_password("securepass123"))
        self.assertTrue(user.password.startswith("argon2$argon2id$"))
        self.assertEqual(user.groups.first().name, "technologist")
        payload.update({"fullName": "Другой Сотрудник", "password": ""})
        updated = self.client.put(f"/api/admin/users/{user.pk}/", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(updated.status_code, 200, updated.content)
        user.refresh_from_db()
        self.assertEqual(user.get_full_name(), "Другой Сотрудник")
        self.assertTrue(user.check_password("securepass123"))
        disabled = self.client.delete(f"/api/admin/users/{user.pk}/")
        self.assertEqual(disabled.status_code, 200)
        user.refresh_from_db()
        self.assertFalse(user.is_active)

    def test_weak_new_password_is_rejected(self):
        response = self.client.post("/api/admin/users/", data=json.dumps({
            "username": "weak_user", "fullName": "Новый Работник", "role": "technologist",
            "password": "12345678"}), content_type="application/json")
        self.assertEqual(response.status_code, 400, response.content)
        self.assertFalse(User.objects.filter(username="weak_user").exists())

    def test_excel_exports_escape_user_supplied_formula(self):
        order = m.Order.objects.first()
        order.number = "=1+1"
        order.save(update_fields=["number"])
        response = self.client.get("/api/legacy/report/export/?kind=production")
        book = load_workbook(io.BytesIO(response.content))
        escaped = [row[0] for row in book.active.iter_rows(min_row=2) if row[0].value == "'=1+1"]
        self.assertEqual(len(escaped), 1)
        self.assertEqual(escaped[0].data_type, "s")
        response = self.client.get("/reports/export/")
        book = load_workbook(io.BytesIO(b"".join(response.streaming_content)))
        self.assertIn("'=1+1", [row[0].value for row in book.active.iter_rows(min_row=2)])

    def test_seeded_business_relationships_are_consistent(self):
        for model in (m.Order, m.OrderItem, m.Program, m.ProgramLine, m.Job,
                      m.ProductionEvent, m.QualityCheck, m.MaterialRequirement,
                      m.RouteOperation, m.Employee):
            for obj in model.objects.all():
                with self.subTest(model=model.__name__, pk=obj.pk):
                    obj.full_clean()
        for job in m.Job.objects.select_related("operation", "program_line__order_item"):
            self.assertEqual(job.operation.product_id, job.program_line.order_item.product_id)
        for event in m.ProductionEvent.objects.select_related("job__program_line__order_item"):
            self.assertEqual(event.order_id, event.job.program_line.order_item.order_id)

    def test_backup_actions_require_admin_and_csrf(self):
        anonymous = Client(HTTP_HOST="127.0.0.1")
        self.assertEqual(anonymous.get("/api/admin/backup/").status_code, 302)
        restricted = Client(HTTP_HOST="127.0.0.1")
        restricted.login(username="user", password="user")
        self.assertEqual(restricted.get("/api/admin/backup/").status_code, 403)
        self.assertEqual(restricted.delete("/api/admin/backup/backup-20261021-100000-12345678.sql/").status_code, 403)
        strict = Client(HTTP_HOST="127.0.0.1", enforce_csrf_checks=True)
        strict.login(username="admin", password="admin")
        self.assertEqual(strict.post("/api/admin/backup/").status_code, 403)
        self.assertEqual(strict.delete("/api/admin/backup/backup-20261021-100000-12345678.sql/").status_code, 403)

    def test_plan_preview_and_approval_use_orders(self):
        preview = self.client.get("/api/legacy/plans/", {"kind": "annual", "year": 2026})
        self.assertEqual(preview.status_code, 200, preview.content)
        self.assertEqual(len(preview.json()["items"]), 3)
        approved = self.client.post("/api/legacy/plans/", data=json.dumps({"kind": "annual", "year": 2026}), content_type="application/json")
        self.assertEqual(approved.status_code, 200, approved.content)
        program = m.Program.objects.get(name="Годовой план 2026")
        self.assertEqual(program.status, "approved")
        self.assertEqual(program.lines.count(), 3)
        repeated = self.client.post("/api/legacy/plans/", data=json.dumps({"kind": "annual", "year": 2026}), content_type="application/json")
        self.assertEqual(repeated.status_code, 200)
        self.assertEqual(program.lines.count(), 3)

    def test_plan_b_preview_distributes_remaining_hours_by_operation(self):
        order = m.Order.objects.get(number="З-2026-041")
        quarter = (order.due_date.month - 1) // 3 + 1
        response = self.client.get("/api/legacy/plans/", {
            "kind": "b", "year": order.due_date.year, "quarter": quarter})
        self.assertEqual(response.status_code, 200, response.content)
        for item in response.json()["items"]:
            values = [Decimal(value) for value in item["remaining_by_category"]]
            self.assertEqual(len(values), 8)
            self.assertLessEqual(sum(values), Decimal(item["hours"]))
        self.assertTrue(any(Decimal(item["remaining_by_category"][7]) > 0
            for item in response.json()["items"]))

    def test_approved_plan_rejects_stale_order_quantities(self):
        payload = {"kind": "annual", "year": 2026}
        approved = self.client.post("/api/legacy/plans/", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(approved.status_code, 200, approved.content)
        line = m.ProgramLine.objects.get(program__name="Годовой план 2026", order_item__order__number="З-2026-041")
        item = line.order_item
        item.quantity += 1
        item.full_clean()
        item.save(update_fields=["quantity"])
        repeated = self.client.post("/api/legacy/plans/", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(repeated.status_code, 400, repeated.content)
        self.assertIn("новую редакцию", repeated.json()["error"])
        line.refresh_from_db()
        self.assertEqual(line.quantity, 12)

    def test_legacy_report_uses_database_and_exports_files(self):
        response = self.client.get("/api/legacy/report/", {"kind": "plan-a"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()["rows"]), 3)
        xlsx = self.client.get("/api/legacy/report/export/", {"kind": "plan-a"})
        self.assertEqual(xlsx.status_code, 200)
        self.assertEqual(xlsx["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        archive = self.client.get("/api/legacy/report/export/", {"kind": "all"})
        with zipfile.ZipFile(io.BytesIO(archive.content)) as files:
            self.assertEqual(len(files.namelist()), 5)

    def test_report_types_return_their_own_columns_and_plan_b_lines(self):
        order = m.Order.objects.get(number="З-2026-041")
        quarter = (order.due_date.month - 1) // 3 + 1
        approved = self.client.post("/api/legacy/plans/", data=json.dumps({
            "kind": "b", "year": order.due_date.year, "quarter": quarter,
        }), content_type="application/json")
        self.assertEqual(approved.status_code, 200, approved.content)
        plan_b = self.client.get("/api/legacy/report/", {"kind": "plan-b"}).json()
        balance = self.client.get("/api/legacy/report/", {"kind": "balance"}).json()
        production = self.client.get("/api/legacy/report/", {"kind": "production"}).json()
        self.assertIn("Рабочий центр", plan_b["headers"])
        self.assertTrue(any(row[0] == order.number for row in plan_b["rows"]))
        self.assertIn("Остаток, шт.", balance["headers"])
        self.assertIn("Готовность, %", production["headers"])
        self.assertEqual(len(plan_b["rows"][0]), len(plan_b["headers"]))

    def test_material_catalog_reads_postgresql_stock(self):
        material = m.Material.objects.first()
        material.stock_quantity = Decimal("17.500")
        material.reserved_quantity = Decimal("2.500")
        material.save(update_fields=["stock_quantity", "reserved_quantity"])
        response = self.client.get("/api/legacy/materials/")
        self.assertEqual(response.status_code, 200)
        row = next(item for item in response.json()["materials"] if item["code"] == material.code)
        self.assertEqual(row["available"], "15.000")

    def test_home_metrics_use_orders_and_work_centers(self):
        response = self.client.get("/api/legacy/dashboard/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["orders"], 3)
        self.assertEqual(response.json()["quantity"], sum(item.quantity for item in m.OrderItem.objects.all()))
        self.assertEqual(response.json()["workplaces"], m.WorkCenter.objects.filter(archived=False).count())

    def test_database_settings_validate_before_save_and_require_admin(self):
        current = db_admin.active_config()
        invalid = {**current, "port": "0"}
        response = self.client.post("/api/admin/database/config/", data=json.dumps(invalid), content_type="application/json")
        self.assertEqual(response.status_code, 400)
        with patch.object(db_admin, "save_config") as saved:
            unchanged = self.client.post("/api/admin/database/config/", data=json.dumps({**current, "password": ""}),
                content_type="application/json")
            self.assertEqual(unchanged.status_code, 200, unchanged.content)
            self.assertFalse(unchanged.json()["restart"])
            saved.assert_not_called()
        changed_config = {**current, "host": "db.example.test"}
        with patch.object(db_admin, "checked_config", return_value=changed_config), patch.object(db_admin, "save_config") as saved:
            changed = self.client.post("/api/admin/database/config/", data=json.dumps(changed_config),
                content_type="application/json")
            self.assertEqual(changed.status_code, 200, changed.content)
            self.assertTrue(changed.json()["restart"])
            saved.assert_called_once_with(changed_config)
        self.client.logout()
        self.client.login(username="user", password="user")
        self.assertEqual(self.client.post("/api/admin/database/config/", data=json.dumps(current),
            content_type="application/json").status_code, 403)

    def test_release_health_endpoints(self):
        self.client.logout()
        self.assertEqual(self.client.get("/health/live/").json(), {"ok": True})
        self.assertEqual(self.client.get("/health/ready/").json(), {"ok": True})

    def test_backup_export_sql_delete_and_restore_requires_supervisor(self):
        exported = self.client.post("/api/admin/backup/export/")
        self.assertEqual(exported.status_code, 200)
        sql = b"".join(exported.streaming_content)
        self.assertIn(b"PostgreSQL database dump", sql[:300])
        self.assertIn(b"django_migrations", sql)
        self.assertTrue(exported["Content-Disposition"].endswith('.sql"'))
        archive = self.client.post("/api/admin/backup/export/?format=dump")
        self.assertEqual(archive.status_code, 200)
        dump = b"".join(archive.streaming_content)
        self.assertTrue(dump.startswith(b"PGDMP"))
        with tempfile.TemporaryDirectory() as temp, patch.object(db_admin, "BACKUP_DIR", Path(temp)):
            created = self.client.post("/api/admin/backup/")
            self.assertEqual(created.status_code, 200, created.content)
            name = created.json()["backups"][0]["name"]
            self.assertTrue(name.endswith(".sql"))
            downloaded = self.client.get(f"/api/admin/backup/{name}/")
            self.assertEqual(downloaded.status_code, 200)
            self.assertIn(b"PostgreSQL database dump", b"".join(downloaded.streaming_content)[:300])
            deleted = self.client.delete(f"/api/admin/backup/{name}/")
            self.assertEqual(deleted.status_code, 200, deleted.content)
            self.assertFalse((Path(temp) / name).exists())
            self.assertEqual(self.client.get(f"/api/admin/backup/{name}/").status_code, 404)
            self.assertEqual(self.client.delete("/api/admin/backup/not-a-backup.sql/").status_code, 404)
        with tempfile.TemporaryDirectory() as temp, patch.object(db_admin, "RUNTIME_DIR", Path(temp)):
            upload = SimpleUploadedFile("database.dump", dump, content_type="application/octet-stream")
            unsupervised = self.client.post("/api/admin/backup/restore/", {"file": upload})
            self.assertEqual(unsupervised.status_code, 503)
            (Path(temp) / "supervisor.json").write_text(json.dumps({"time": time.time(),
                "database": db_admin.active_config()["database"]}), encoding="utf-8")
            upload = SimpleUploadedFile("database.dump", dump, content_type="application/octet-stream")
            queued = self.client.post("/api/admin/backup/restore/", {"file": upload})
            self.assertEqual(queued.status_code, 202, queued.content)
            self.assertTrue((Path(temp) / "restore-request.json").exists())

    def test_process_actions_are_database_backed(self):
        listing = self.client.get("/api/legacy/processes/")
        self.assertEqual(listing.status_code, 200)
        self.assertEqual(len(listing.json()["processes"]), 3)
        route = m.RouteOperation.objects.filter(job__isnull=False).first()
        approved = self.client.put(f"/api/legacy/processes/{route.pk}/")
        self.assertEqual(approved.status_code, 200)
        route.refresh_from_db()
        self.assertTrue(route.approved)
        protected = self.client.delete(f"/api/legacy/processes/{route.pk}/")
        self.assertEqual(protected.status_code, 409)

    def test_process_editor_creates_and_updates_route_without_changing_page_design(self):
        listing = self.client.get("/api/legacy/processes/").json()
        self.assertTrue(listing["products"])
        self.assertTrue(listing["centers"])
        payload = {"product_id": listing["products"][0]["id"],
            "work_center_id": listing["centers"][0]["id"], "number": 99,
            "name": "Контрольная операция", "kind": "machining", "setup_hours": "1.25",
            "piece_hours": "0.5", "quality_required": True,
            "shop_code": "01", "section_code": "02", "workplace_code": "03",
            "operation_text": "Проверить поверхность", "tooling": "Калибр; микрометр",
            "worker_count": 2, "simultaneous_quantity": 1}
        created = self.client.post("/api/legacy/processes/", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(created.status_code, 200, created.content)
        pk = created.json()["id"]
        payload["piece_hours"] = "0.75"
        updated = self.client.put(f"/api/legacy/processes/{pk}/", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(updated.status_code, 200, updated.content)
        self.assertEqual(m.RouteOperation.objects.get(pk=pk).piece_hours, Decimal("0.750"))
        self.assertEqual(m.RouteOperation.objects.get(pk=pk).shop_code, "01")
        self.assertEqual(m.RouteOperation.objects.get(pk=pk).tooling, "Калибр; микрометр")

    def test_deviation_create_edit_delete_in_postgresql(self):
        payload = {"order": "З-2026-041", "date": "2026-10-14", "hours": "2.5", "reason": "Отсутствие инструмента/оснастки"}
        created = self.client.post("/api/legacy/deviations/", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(created.status_code, 200, created.content)
        pk = created.json()["deviation"]["id"]
        item = m.Downtime.objects.get(pk=pk)
        self.assertEqual(item.hours, Decimal("2.5"))
        self.assertEqual(item.reason, "Прочее")
        payload["hours"] = "3.5"
        updated = self.client.put(f"/api/legacy/deviations/{pk}/", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(updated.status_code, 200, updated.content)
        item.refresh_from_db()
        self.assertEqual(item.hours, Decimal("3.5"))
        self.assertEqual(self.client.delete(f"/api/legacy/deviations/{pk}/").status_code, 200)
        self.assertFalse(m.Downtime.objects.filter(pk=pk).exists())

    def test_login_redirect_rejects_external_destinations(self):
        for destination in ("https://outside.example/steal", "//outside.example/steal"):
            self.client.logout()
            response = self.client.post(f"/login/?next={destination}",
                {"username": "admin", "password": "admin"})
            self.assertEqual(response.status_code, 302)
            self.assertEqual(response["Location"], "/")
        self.client.logout()
        response = self.client.post("/login/?next=/orders.html",
            {"username": "admin", "password": "admin"})
        self.assertEqual(response["Location"], "/orders.html")

    def test_unassigned_and_ambiguous_roles_cannot_enter_application(self):
        extra = User.objects.create_user("no_role", password="A long unique password 2026")
        self.client.logout()
        response = self.client.post("/login/", {"username": "no_role", "password": "A long unique password 2026"})
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("_auth_user_id", self.client.session)
        extra.groups.set([Group.objects.get(name="planner"), Group.objects.get(name="technologist")])
        self.client.force_login(extra)
        self.assertEqual(self.client.get("/").status_code, 403)
        self.assertEqual(self.client.get("/api/legacy/orders/").status_code, 403)

    def test_repeated_login_failures_are_limited(self):
        cache.clear()
        self.client.logout()
        for _ in range(5):
            self.assertEqual(self.client.post("/login/", {"username": "admin", "password": "wrong"}).status_code, 200)
        blocked = self.client.post("/login/", {"username": "admin", "password": "admin"})
        self.assertEqual(blocked.status_code, 429)
        cache.clear()
        self.assertEqual(self.client.post("/login/", {"username": "admin", "password": "admin"}).status_code, 302)

    def test_direct_api_and_pdf_urls_obey_role_permissions(self):
        job = m.Job.objects.first()
        self.client.logout()
        self.client.login(username="user", password="user")
        self.assertEqual(self.client.get("/api/legacy/report/").status_code, 403)
        self.assertEqual(self.client.get("/api/legacy/report/export/").status_code, 403)
        self.assertEqual(self.client.get("/api/legacy/deviations/").status_code, 403)
        self.assertEqual(self.client.get("/reports/").status_code, 403)
        self.assertEqual(self.client.get(f"/jobs/{job.pk}/print/").status_code, 403)
        self.assertEqual(self.client.get("/api/legacy/materials/").status_code, 200)
        self.client.logout()
        self.client.login(username="planner", password="planner")
        self.assertEqual(self.client.get("/api/legacy/materials/").status_code, 403)
        self.assertEqual(self.client.get("/api/legacy/processes/").status_code, 200)

    def test_backup_export_uses_csrf_protected_post(self):
        self.assertEqual(self.client.get("/api/admin/backup/export/").status_code, 405)
        strict = Client(HTTP_HOST="127.0.0.1", enforce_csrf_checks=True)
        strict.login(username="admin", password="admin")
        self.assertEqual(strict.post("/api/admin/backup/export/").status_code, 403)

    def test_authenticated_pages_are_not_browser_cached(self):
        response = self.client.get("/orders.html")
        self.assertEqual(response.status_code, 200)
        self.assertIn("no-store", response["Cache-Control"])

    def test_xlsx_import_rejects_excessive_uncompressed_size(self):
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as book:
            book.writestr("xl/worksheets/sheet1.xml", b"X" * (51 * 1024 * 1024))
        upload = SimpleUploadedFile("huge.xlsx", archive.getvalue())
        response = self.client.post("/materials/import/", {"file": upload}, HTTP_ACCEPT="application/json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("превышает", response.json()["error"])

    def test_estd_form_one_uses_linked_material_route_and_order_quantity(self):
        product = m.Product.objects.get(name="Винт М12")
        item = m.OrderItem.objects.get(product=product)
        product.material_specification = "Демонстрационный материал"
        product.material_code = "ДЕМО-45"
        product.part_mass_kg = Decimal("0.500")
        product.material_consumption_kg = Decimal("0.800")
        product.blank_type = "Прокат"
        product.blank_dimensions = "Круг 30 x 150"
        product.blank_mass_kg = Decimal("0.800")
        product.full_clean()
        product.save()
        operation = product.route.get()
        operation.shop_code = "01"
        operation.section_code = "02"
        operation.workplace_code = "03"
        operation.operation_text = "Обточить поверхность до размера по чертежу"
        operation.tooling = "Патрон; резец; штангенциркуль"
        operation.full_clean()
        operation.save()
        card = route_card_data(product, item)
        self.assertEqual(card["gaps"], [])
        self.assertEqual(card["material"]["kim"], "0,625")
        self.assertEqual(card["operations"][0]["a"][:4], ["01", "02", "03", "010"])
        self.assertEqual(card["operations"][0]["b"][6], item.quantity)
        self.assertEqual(card["operations"][0]["b"][7], "45,00")
        self.assertEqual(card["operations"][0]["b"][8], "18,00")
        self.assertTrue(render_route_card(card).getvalue().startswith(b"%PDF"))
        pdf = self.client.get(f"/products/{product.pk}/route.pdf?item={item.pk}")
        self.assertEqual(pdf.status_code, 200)
        self.assertIn("route-gost-3.1118", pdf["Content-Disposition"])
        other_item = m.OrderItem.objects.exclude(product=product).first()
        self.assertEqual(self.client.get(f"/products/{product.pk}/route.pdf?item={other_item.pk}").status_code, 404)

    def test_estd_product_editor_validates_and_saves_material(self):
        product = m.Product.objects.get(name="Винт М12")
        payload = {"material_specification": "Учебный материал", "material_code": "ДЕМО",
            "part_mass_kg": "0.500", "material_consumption_kg": "0.800",
            "blank_type": "Прокат", "blank_dimensions": "Круг 30 x 150",
            "blank_mass_kg": "0.800", "parts_per_blank": 2}
        endpoint = f"/api/legacy/products/{product.pk}/estd/"
        saved = self.client.put(endpoint, data=json.dumps(payload), content_type="application/json")
        self.assertEqual(saved.status_code, 200, saved.content)
        product.refresh_from_db()
        self.assertEqual(product.parts_per_blank, 2)
        self.assertEqual(product.part_mass_kg, Decimal("0.500"))
        self.assertTrue(m.AuditLog.objects.filter(model_name="Product", object_id=product.pk,
            reason="Паспорт материала маршрутной карты ЕСТД").exists())
        invalid = self.client.put(endpoint, data=json.dumps({**payload, "part_mass_kg": "-1"}),
            content_type="application/json")
        self.assertEqual(invalid.status_code, 400)
        product.refresh_from_db()
        self.assertEqual(product.part_mass_kg, Decimal("0.500"))
        self.client.logout()
        self.client.login(username="planner", password="planner")
        self.assertEqual(self.client.get(endpoint).status_code, 200)
        self.assertEqual(self.client.put(endpoint, data=json.dumps(payload),
            content_type="application/json").status_code, 403)

    def test_assembly_route_is_marked_as_needing_separate_estd_form(self):
        product = m.Product.objects.get(name="Патрон станка")
        card = route_card_data(product)
        self.assertIn("для сборочного процесса требуется отдельная форма 2", card["gaps"])


class SampleImportTests(TestCase):
    def test_installer_distinguishes_empty_sample_and_existing_database(self):
        output = io.StringIO()
        call_command("sample_data_state", stdout=output)
        self.assertEqual(output.getvalue().strip(), "empty")
        call_command("import_sample_data", verbosity=0)
        output = io.StringIO()
        call_command("sample_data_state", stdout=output)
        self.assertEqual(output.getvalue().strip(), "sample")
        output = io.StringIO()
        call_command("sample_data_state", usernames=True, stdout=output)
        self.assertEqual(set(output.getvalue().splitlines()), {"user", "planner", "manager", "supervisor"})
        User.objects.create_user("extra_user", password="Secure-random-test-password-2026")
        output = io.StringIO()
        call_command("sample_data_state", stdout=output)
        self.assertEqual(output.getvalue().strip(), "populated")

    def test_fixture_import_and_first_admin_do_not_publish_passwords(self):
        call_command("import_sample_data", verbosity=0)
        self.assertEqual(m.Order.objects.count(), 3)
        self.assertEqual(m.RouteOperation.objects.count(), 3)
        self.assertFalse(User.objects.get(username="admin").has_usable_password())
        self.assertFalse(User.objects.get(username="user").is_active)
        with self.assertRaises(CommandError):
            call_command("import_sample_data", verbosity=0)
        with patch.dict("os.environ", {"AVT_ADMIN_PASSWORD": "A-new-long-random-password-2026"}):
            call_command("provision_admin", verbosity=0)
        admin = User.objects.get(username="admin")
        self.assertTrue(admin.is_active and admin.is_superuser)
        self.assertTrue(admin.check_password("A-new-long-random-password-2026"))
        with patch.dict("os.environ", {"AVT_ADMIN_PASSWORD": "A-different-long-random-password"}):
            call_command("provision_admin", verbosity=0)
        admin.refresh_from_db()
        self.assertTrue(admin.check_password("A-new-long-random-password-2026"))

    def test_reset_password_activates_only_requested_existing_user(self):
        call_command("import_sample_data", verbosity=0)
        with patch.dict("os.environ", {"AVT_NEW_PASSWORD": "Secure-random-test-password-2026"}):
            call_command("reset_user_password", "user", activate=True, verbosity=0)
        user = User.objects.get(username="user")
        self.assertTrue(user.is_active)
        self.assertTrue(user.check_password("Secure-random-test-password-2026"))
        self.assertTrue(user.password.startswith("argon2$argon2id$"))
        self.assertFalse(User.objects.get(username="planner").is_active)
        with self.assertRaises(CommandError):
            call_command("reset_user_password", "missing", verbosity=0)


class ReleaseSurfaceTests(TestCase):
    def test_user_form_creates_role_in_empty_database(self):
        form = UserForm(data={"username": "new_manager", "first_name": "Иван",
            "last_name": "Тестовый", "email": "", "is_active": True,
            "role": "manager", "password": "Secure-random-test-password-2026"})
        self.assertTrue(form.is_valid(), form.errors)
        user = form.save()
        self.assertEqual(list(user.groups.values_list("name", flat=True)), ["manager"])
        self.assertTrue(user.check_password("Secure-random-test-password-2026"))

    def test_login_has_favicon_and_no_published_test_passwords(self):
        response = Client(HTTP_HOST="127.0.0.1").get("/login/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "favicon.svg")
        self.assertNotContains(response, "Демо-доступ")
        self.assertNotContains(response, "admin / admin")
        self.assertNotContains(response, "user / user")

import csv
import hashlib
import io
from html import escape
import json
import itertools
import logging
import zipfile
from datetime import timedelta
from decimal import Decimal
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.core.cache import cache
from django.core.exceptions import PermissionDenied
from django.core.exceptions import ValidationError
from django.core.serializers.json import DjangoJSONEncoder
from django.db import IntegrityError, transaction
from django.db.models import F, Q, Sum
from django.http import FileResponse, Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.utils import simpleSplit
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from .forms import REGISTRY, LoginForm, UserForm, form_for
from .export_security import spreadsheet_value
from . import models
from .scheduling import assign_worker
from .estd import _font as cyrillic_pdf_font, render_route_card, route_card_data

logger = logging.getLogger(__name__)


ROLE_ACCESS = {
    "admin": set(REGISTRY),
    "manager": set(REGISTRY),
    "planner": {"orders", "items", "programs", "lines", "jobs", "events", "quality", "materials", "requirements", "downtime", "products", "routes", "customers", "departments", "centers", "equipment", "employees"},
    "technologist": {"orders", "items", "products", "routes", "cutting-norms", "materials", "customers", "departments", "centers", "equipment"},
    "normalizer": {"orders", "items", "programs", "lines", "jobs", "products", "routes", "cutting-norms", "centers", "employees"},
    "supervisor": {"orders", "items", "jobs", "events", "quality", "materials", "requirements", "downtime", "products", "routes", "centers", "equipment", "employees"},
    "economist": {"orders", "items", "programs", "lines", "jobs", "events", "materials", "requirements", "products"},
}
ROLE_WRITE = {
    "admin": set(REGISTRY), "manager": set(REGISTRY),
    "planner": {"orders", "items", "programs", "lines", "jobs", "requirements", "downtime"},
    "technologist": {"orders", "items", "products", "routes", "materials"},
    "normalizer": {"cutting-norms"},
    "supervisor": {"jobs", "events", "quality", "downtime"},
    "economist": set(),
}
REPORT_ROLES = {"admin", "manager", "planner", "supervisor", "normalizer", "economist"}


def role(user):
    if user.is_superuser:
        return "admin"
    groups = list(user.groups.values_list("name", flat=True))
    return groups[0] if len(groups) == 1 and groups[0] in ROLE_ACCESS else "unassigned"


def access(request, key, write=False):
    if key not in REGISTRY:
        raise Http404
    permitted = ROLE_WRITE if write else ROLE_ACCESS
    if not request.user.is_superuser and key not in permitted.get(role(request.user), set()):
        raise PermissionDenied


def context(request, **extra):
    user_role = role(request.user)
    nav = [(key, title) for key, (_, title, _) in REGISTRY.items() if request.user.is_superuser or key in ROLE_ACCESS.get(user_role, set())]
    return {"nav": nav, "role": user_role, **extra}


def sign_in(request):
    if request.user.is_authenticated:
        return redirect("legacy_home")
    from .legacy_ui import render_login
    form = LoginForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        username = form.cleaned_data["username"]
        identity = hashlib.sha256(username.casefold().encode("utf-8")).hexdigest()
        remote = hashlib.sha256((request.META.get("REMOTE_ADDR") or "unknown").encode("utf-8")).hexdigest()
        user_key, ip_key = f"login-fail-user:{identity}", f"login-fail-ip:{remote}"
        if cache.get(user_key, 0) >= 5 or cache.get(ip_key, 0) >= 30:
            form.add_error(None, "Слишком много попыток входа. Повторите через 15 минут")
            response = render_login(request, form)
            response.status_code = 429
            return response
        user = authenticate(request, username=username, password=form.cleaned_data["password"])
        if user and user.is_active and role(user) != "unassigned":
            policy = models.SitePolicy.objects.filter(pk=1, maintenance=True).only("maintenance_owner_id").first()
            if policy and policy.maintenance_owner_id != user.pk:
                form.add_error(None, "Вход временно приостановлен: выполняется техническое обслуживание.")
                return render_login(request, form)
            cache.delete(user_key)
            login(request, user)
            models.AuditLog.objects.create(user=user, model_name="Session", object_id=user.pk, action="login")
            next_url = request.POST.get("next") or request.GET.get("next", "")
            if next_url and url_has_allowed_host_and_scheme(next_url, allowed_hosts={request.get_host()},
                                                              require_https=request.is_secure()):
                return redirect(next_url)
            return redirect("legacy_home")
        for key in (user_key, ip_key):
            cache.add(key, 0, timeout=15 * 60)
            cache.incr(key)
        form.add_error(None, "Неверный логин или пароль")
    return render_login(request, form)


@login_required
@require_POST
def sign_out(request):
    models.AuditLog.objects.create(user=request.user, model_name="Session", object_id=request.user.pk, action="logout")
    logout(request)
    return redirect("login")


@login_required
def dashboard(request):
    orders = models.Order.objects.filter(archived=False)
    jobs = models.Job.objects.all()
    today = timezone.localdate()
    planned = models.ProgramLine.objects.filter(program__status__in=["approved", "revised"]).aggregate(v=Sum("quantity"))["v"] or 0
    accepted = models.ProductionEvent.objects.filter(kind="accepted").aggregate(v=Sum("quantity"))["v"] or 0
    recent_orders = list(orders.select_related("customer", "department")[:8])
    at_risk_count = sum(1 for order in orders if order.is_at_risk)
    return render(request, "production/dashboard.html", context(request,
        active_orders=orders.exclude(status__in=["done", "closed"]).count(),
        active_jobs=jobs.filter(status__in=["work", "paused", "quality"]).count(),
        overdue_orders=orders.filter(due_date__lt=today).exclude(status__in=["done", "closed"]).count(),
        overdue_jobs=jobs.filter(planned_end__lt=today).exclude(status__in=["done", "accepted"]).count(),
        at_risk=at_risk_count,
        planned=planned, accepted=accepted,
        completion=round(accepted * 100 / planned) if planned else 0,
        recent_orders=recent_orders))


def queryset_for(key):
    model = REGISTRY[key][0]
    qs = model.objects.all()
    if hasattr(model, "archived"):
        qs = qs.filter(archived=False)
    return qs


def display_columns(key):
    preferred = {
        "orders": ["number", "customer", "department", "start_date", "due_date", "forecast_date", "priority", "status", "progress"],
        "items": ["order", "product", "batch", "quantity"],
        "programs": ["name", "year", "quarter", "version", "status"],
        "lines": ["program", "order_item", "work_center", "quantity", "due_date", "forecast_date", "planned_hours"],
        "jobs": ["program_line", "operation", "employee", "quantity", "actual_quantity", "planned_end", "priority", "status", "planned_hours"],
        "events": ["order", "job", "kind", "quantity", "date", "document"],
        "quality": ["job", "result", "quantity", "defect", "date", "inspector"],
        "materials": ["code", "name", "unit", "stock_quantity", "reserved_quantity", "available_quantity"],
        "requirements": ["order_item", "material", "quantity_per_item", "required", "reserved", "deficit"],
        "routes": ["product", "number", "name", "kind", "work_center", "setup_hours", "piece_hours", "approved"],
        "cutting-norms": ["operation", "material", "tool", "source_document", "spindle_rpm", "feed_mm_min", "machine_hours", "approved"],
    }
    model = REGISTRY[key][0]
    names = preferred.get(key) or [f.name for f in model._meta.fields if f.name not in ("id", "archived", "created_at", "updated_at")][:7]
    labels = []
    for name in names:
        try:
            label = str(model._meta.get_field(name).verbose_name)
        except Exception:
            label = {"progress": "Готовность, %", "planned_hours": "Норма, ч", "required": "Потребность", "deficit": "Дефицит", "available_quantity": "Доступно", "priority": "Приоритет", "forecast_date": "Прогноз выпуска"}.get(name, name)
        labels.append((name, label))
    return labels


@login_required
def records(request, key):
    access(request, key)
    qs = queryset_for(key)
    q = request.GET.get("q", "").strip()
    if q:
        text_fields = [f.name for f in REGISTRY[key][0]._meta.fields if f.get_internal_type() in ("CharField", "TextField")]
        if text_fields:
            condition = Q()
            for field in text_fields:
                condition |= Q(**{f"{field}__icontains": q})
            qs = qs.filter(condition)
    if key == "orders":
        for param in ("status", "priority", "department"):
            if request.GET.get(param):
                qs = qs.filter(**{param: request.GET[param]})
        if request.GET.get("due_before"):
            qs = qs.filter(due_date__lte=request.GET["due_before"])
    if key == "jobs" and request.GET.get("status"):
        qs = qs.filter(status=request.GET["status"])
    if key == "jobs":
        if request.GET.get("work_center"):
            qs = qs.filter(operation__work_center_id=request.GET["work_center"])
        if request.GET.get("employee"):
            qs = qs.filter(employee_id=request.GET["employee"])
        if request.GET.get("due_before"):
            qs = qs.filter(planned_end__lte=request.GET["due_before"])
        qs = qs.order_by("-program_line__order_item__order__priority", "planned_end")
    columns = display_columns(key)
    rows = []
    for obj in qs[:500]:
        cells = []
        for field, _ in columns:
            value = getattr(obj, field)
            if field == "progress":
                value = f"{value}%"
            elif field in ("status", "kind", "result"):
                method = getattr(obj, f"get_{field}_display", None)
                value = method() if method else value
            cells.append(value)
        rows.append((obj, cells))
    return render(request, "production/records.html", context(request,
        key=key, title=REGISTRY[key][1], singular=REGISTRY[key][2], columns=columns,
        rows=rows, q=q, can_write=request.user.is_superuser or key in ROLE_WRITE.get(role(request.user), set()),
        statuses=models.Order.STATUSES if key == "orders" else models.Job.STATUSES if key == "jobs" else [],
        departments=models.Department.objects.filter(archived=False) if key == "orders" else [],
        products=models.Product.objects.filter(archived=False) if key == "routes" else [],
        centers=models.WorkCenter.objects.filter(archived=False) if key == "jobs" else [],
        employees=models.Employee.objects.filter(archived=False) if key == "jobs" else []))


def snapshot(obj):
    if obj is None:
        return {}
    return json.loads(json.dumps({f.name: getattr(obj, f.attname) for f in obj._meta.fields if f.name != "id"}, cls=DjangoJSONEncoder))


@login_required
def record_form(request, key, pk=None):
    access(request, key, write=True)
    model = REGISTRY[key][0]
    instance = get_object_or_404(model, pk=pk) if pk else None
    if key == "requirements" and instance and instance.reserved:
        raise PermissionDenied("Сначала снимите резерв перед изменением потребности.")
    if instance and key == "events" and role(request.user) not in ("admin", "manager") and not request.user.is_superuser:
        raise PermissionDenied
    posted_program = request.POST.get("program") if request.method == "POST" else None
    if key == "lines" and ((instance and instance.program.status != "draft") or (posted_program and models.Program.objects.filter(pk=posted_program, status__in=["approved", "revised", "closed"]).exists())):
        raise PermissionDenied
    before = snapshot(instance)
    form = form_for(key, request.POST or None, instance=instance)
    if key == "routes" and role(request.user) not in ("admin", "manager"):
        # Технолог готовит маршрут, а утверждает его начальник отдела.
        form.fields.pop("approved", None)
    if request.method == "POST" and form.is_valid():
        if instance and not request.POST.get("change_reason", "").strip():
            form.add_error(None, "Укажите основание корректировки")
        else:
            return save_record(request, key, model, instance, before, form)
    template = "production/form_fragment.html" if request.headers.get("x-requested-with") == "XMLHttpRequest" else "production/form_page.html"
    result = render(request, template, context(request, key=key, form=form, object=instance,
        title=("Изменить " if instance else "Добавить ") + REGISTRY[key][2].lower()))
    if request.headers.get("x-requested-with") == "XMLHttpRequest" and request.method == "POST":
        result.status_code = 400
    return result


def save_record(request, key, model, instance, before, form):
    """Persist one validated form together with its audit trail."""
    try:
        with transaction.atomic():
            obj = form.save(commit=False)
            if key == "routes" and role(request.user) not in ("admin", "manager"):
                obj.approved = False
            if key == "events" and not instance:
                obj.actor = request.user
            obj.full_clean()
            if key == "programs" and obj.status in ("approved", "revised") and (not obj.pk or not obj.lines.exists()):
                raise ValidationError("Нельзя утвердить программу без строк плана выпуска.")
            obj.save()
            if key == "programs" and obj.status in ("approved", "revised"):
                models.Program.objects.filter(name=obj.name, status__in=["approved", "revised"]).exclude(pk=obj.pk).update(status="closed")
            models.AuditLog.objects.create(user=request.user, model_name=model.__name__, object_id=obj.pk,
                action="update" if instance else "create", reason=request.POST.get("change_reason", ""),
                before=before, after=snapshot(obj))
        messages.success(request, "Запись сохранена")
        if request.headers.get("x-requested-with") == "XMLHttpRequest":
            return JsonResponse({"ok": True, "redirect": reverse("records", args=[key])})
        return redirect("records", key=key)
    except (ValidationError, IntegrityError) as exc:
        form.add_error(None, str(exc))
    template = "production/form_fragment.html" if request.headers.get("x-requested-with") == "XMLHttpRequest" else "production/form_page.html"
    result = render(request, template, context(request, key=key, form=form, object=instance,
        title=("Изменить " if instance else "Добавить ") + REGISTRY[key][2].lower()))
    if request.headers.get("x-requested-with") == "XMLHttpRequest":
        result.status_code = 400
    return result


@login_required
@require_POST
def archive_record(request, key, pk):
    access(request, key, write=True)
    model = REGISTRY[key][0]
    obj = get_object_or_404(model, pk=pk)
    if key == "routes" and obj.approved and role(request.user) not in ("admin", "manager"):
        raise PermissionDenied("Утверждённый маршрут может удалить только руководитель.")
    if key == "events" and role(request.user) not in ("admin", "manager") and not request.user.is_superuser:
        raise PermissionDenied
    if key == "lines" and obj.program.status != "draft":
        raise PermissionDenied
    if key == "requirements" and obj.reserved:
        messages.error(request, "Сначала снимите резерв материала")
        return redirect("records", key=key)
    reason = request.POST.get("reason", "").strip()
    if not reason:
        messages.error(request, "Укажите основание изменения")
        return redirect("records", key=key)
    before = snapshot(obj)
    with transaction.atomic():
        if hasattr(obj, "archived"):
            obj.archived = True
            obj.save(update_fields=["archived"])
        else:
            try:
                obj.delete()
            except Exception as exc:
                messages.error(request, f"Удаление запрещено: {exc}")
                return redirect("records", key=key)
        models.AuditLog.objects.create(user=request.user, model_name=model.__name__, object_id=pk,
            action="archive" if hasattr(obj, "archived") else "delete", reason=reason, before=before)
    messages.success(request, "Запись архивирована или удалена")
    return redirect("records", key=key)


@login_required
def order_detail(request, pk):
    order = get_object_or_404(models.Order.objects.select_related("customer", "department"), pk=pk)
    return render(request, "production/order_detail.html", context(request, order=order,
        items=order.items.select_related("product"),
        events=order.events.select_related("job", "actor")[:50],
        history=models.AuditLog.objects.filter(model_name="Order", object_id=pk)[:30],
        requirements=models.MaterialRequirement.objects.filter(order_item__order=order).select_related("material", "order_item")))


@login_required
@require_POST
def revise_program(request, pk):
    access(request, "programs", write=True)
    source = get_object_or_404(models.Program, pk=pk)
    if source.status not in ("approved", "revised"):
        messages.error(request, "Новую редакцию можно создать из утвержденной программы")
        return redirect("records", key="programs")
    reason = request.POST.get("reason", "").strip()
    if not reason:
        messages.error(request, "Укажите причину корректировки")
        return redirect("records", key="programs")
    with transaction.atomic():
        latest_version = models.Program.objects.filter(name=source.name).order_by("-version").values_list("version", flat=True).first() or source.version
        new = models.Program.objects.create(name=source.name, year=source.year, quarter=source.quarter,
            version=latest_version + 1, status="draft", correction_reason=reason)
        for line in source.lines.all():
            models.ProgramLine.objects.create(program=new, order_item=line.order_item,
                work_center=line.work_center, quantity=line.quantity, due_date=line.due_date)
        models.AuditLog.objects.create(user=request.user, model_name="Program", object_id=new.pk,
            action="revision", reason=reason, before=snapshot(source), after=snapshot(new))
    messages.success(request, f"Создана версия {new.version}")
    return redirect("records", key="programs")


@login_required
@require_POST
def generate_jobs(request, pk):
    access(request, "jobs", write=True)
    line = get_object_or_404(models.ProgramLine.objects.select_related("program", "order_item__product"), pk=pk)
    if line.program.status not in ("approved", "revised"):
        messages.error(request, "Сначала утвердите производственную программу")
        return redirect("records", key="lines")
    routes = list(line.order_item.product.route.filter(approved=True))
    if not routes:
        messages.error(request, "Для изделия нет утвержденного технологического маршрута")
        return redirect("records", key="lines")
    with transaction.atomic():
        for operation in routes:
            job, created = models.Job.objects.get_or_create(program_line=line, operation=operation,
                defaults={"quantity": line.quantity, "planned_start": line.order_item.order.start_date,
                          "planned_end": line.due_date, "status": "planned"})
            if created:
                models.AuditLog.objects.create(user=request.user, model_name="Job", object_id=job.pk,
                    action="create", reason="Сформировано по утвержденной программе", after=snapshot(job))
    messages.success(request, f"Задания сформированы: {len(routes)} операций маршрута")
    return redirect("records", key="jobs")


@login_required
@require_POST
def assign_job(request, pk):
    access(request, "jobs", write=True)
    job = get_object_or_404(models.Job, pk=pk)
    before = snapshot(job)
    try:
        employee = assign_worker(pk)
    except ValidationError as exc:
        messages.error(request, "; ".join(exc.messages))
        return redirect("records", key="jobs")
    job.refresh_from_db()
    models.AuditLog.objects.create(user=request.user, model_name="Job", object_id=job.pk,
        action="assign", reason="Автоматическая расстановка по квалификации и свободной мощности",
        before=before, after=snapshot(job))
    messages.success(request, f"На задание ПЗ-{pk} назначен {employee.name}")
    return redirect("records", key="jobs")


@login_required
@require_POST
def copy_route(request, pk):
    access(request, "routes", write=True)
    source = get_object_or_404(models.Product, pk=pk)
    target = get_object_or_404(models.Product, pk=request.POST.get("target"))
    reason = request.POST.get("reason", "").strip()
    if not reason or source.pk == target.pk:
        messages.error(request, "Выберите другое изделие и укажите основание копирования")
        return redirect("records", key="routes")
    with transaction.atomic():
        for operation in source.route.all():
            clone, created = models.RouteOperation.objects.get_or_create(product=target, number=operation.number,
                defaults={"name": operation.name, "kind": operation.kind, "definition": operation.definition,
                    "work_center": operation.work_center,
                    "equipment": operation.equipment, "setup_hours": operation.setup_hours,
                    "piece_hours": operation.piece_hours, "quality_required": operation.quality_required,
                    "approved": False})
            if created:
                models.AuditLog.objects.create(user=request.user, model_name="RouteOperation", object_id=clone.pk,
                    action="copy", reason=reason, after=snapshot(clone))
    messages.success(request, "Маршрут скопирован для последующей корректировки")
    return redirect("records", key="routes")


@login_required
def capacity(request):
    start = timezone.localdate()
    if request.GET.get("start"):
        try:
            start = timezone.datetime.fromisoformat(request.GET["start"]).date()
        except ValueError:
            pass
    try:
        days = min(max(int(request.GET.get("days", 30)), 1), 365)
    except ValueError:
        days = 30
    end = start + timedelta(days=days - 1)
    working_days = sum(1 for i in range(days) if (start + timedelta(days=i)).weekday() < 5)
    rows = []
    for center in models.WorkCenter.objects.filter(archived=False).select_related("department"):
        available = center.daily_capacity * working_days
        planned = sum((job.planned_hours for job in models.Job.objects.filter(operation__work_center=center,
            planned_end__gte=start, planned_start__lte=end).select_related("operation")), Decimal("0"))
        actual = models.Job.objects.filter(operation__work_center=center, actual_end__range=(start, end)).aggregate(v=Sum("actual_hours"))["v"] or Decimal("0")
        rows.append((center, available, planned, actual, round(planned * 100 / available) if available else 0))
    employee_rows = []
    for employee in models.Employee.objects.filter(archived=False).select_related("department"):
        available = employee.hours_per_day * working_days
        planned = sum((job.planned_hours for job in models.Job.objects.filter(employee=employee,
            planned_end__gte=start, planned_start__lte=end).select_related("operation")), Decimal("0"))
        actual = models.Job.objects.filter(employee=employee, actual_end__range=(start, end)).aggregate(v=Sum("actual_hours"))["v"] or Decimal("0")
        employee_rows.append((employee, available, planned, actual, round(planned * 100 / available) if available else 0))
    return render(request, "production/capacity.html", context(request, rows=rows, employee_rows=employee_rows,
        start=start.isoformat(), days=days, working_days=working_days))


@login_required
@require_POST
def import_materials(request):
    access(request, "materials", write=True)
    wants_json = "application/json" in request.headers.get("Accept", "")
    upload = request.FILES.get("file")
    if not upload or upload.size > 20 * 1024 * 1024:
        if wants_json:
            return JsonResponse({"ok": False, "error": "Выберите файл CSV, JSON или XLSX размером до 20 МБ"}, status=400)
        messages.error(request, "Выберите файл CSV, JSON или XLSX размером до 20 МБ")
        return redirect("records", key="materials")
    try:
        suffix = upload.name.rsplit(".", 1)[-1].lower()
        if suffix == "json":
            data = json.loads(upload.read().decode("utf-8-sig"))
            rows = data if isinstance(data, list) else data.get("materials", [])
        elif suffix == "csv":
            rows = list(csv.DictReader(io.StringIO(upload.read().decode("utf-8-sig")), delimiter=";"))
            if rows and len(rows[0]) == 1:
                upload.seek(0)
                rows = list(csv.DictReader(io.StringIO(upload.read().decode("utf-8-sig"))))
        elif suffix == "xlsx":
            with zipfile.ZipFile(upload) as archive:
                parts = archive.infolist()
                if (len(parts) > 200 or sum(part.file_size for part in parts) > 50 * 1024 * 1024
                        or any(part.file_size > 20 * 1024 * 1024 for part in parts)):
                    raise ValueError("Книга XLSX превышает допустимый объём данных")
            upload.seek(0)
            book = load_workbook(upload, read_only=True, data_only=True)
            try:
                values = list(itertools.islice(book.active.values, 10_002))
            finally:
                book.close()
            rows = [dict(zip(values[0], row)) for row in values[1:]] if values else []
        else:
            raise ValueError("Поддерживаются CSV, JSON и XLSX")
        if not rows:
            raise ValueError("Файл не содержит строк")
        if len(rows) > 10_000:
            raise ValueError("В одном файле допускается не более 10 000 строк")
        count = 0
        with transaction.atomic():
            for row in rows:
                code = str(row.get("code") or row.get("Код") or row.get("Номенклатурный №") or "").strip()
                name = str(row.get("name") or row.get("Наименование") or row.get("Материал") or "").strip()
                if not code or not name:
                    raise ValueError("В каждой строке нужны код и наименование")
                unit = str(row.get("unit") or row.get("Ед.") or "кг").strip()
                stock = Decimal(str(row.get("stock_quantity") or row.get("Остаток") or 0).replace(",", "."))
                reserved = Decimal(str(row.get("reserved_quantity") or row.get("Зарезервировано") or 0).replace(",", "."))
                material, _ = models.Material.objects.update_or_create(code=code,
                    defaults={"name": name, "unit": unit, "stock_quantity": stock,
                              "reserved_quantity": reserved, "archived": False})
                models.AuditLog.objects.create(user=request.user, model_name="Material", object_id=material.pk,
                    action="import", reason=f"Файл {upload.name}", after=snapshot(material))
                count += 1
        messages.success(request, f"Импортировано материалов: {count}")
        if wants_json:
            return JsonResponse({"ok": True, "count": count})
    except (ValueError, TypeError, UnicodeDecodeError, zipfile.BadZipFile, ValidationError) as exc:
        messages.error(request, f"Ошибка импорта: {exc}")
        if wants_json:
            return JsonResponse({"ok": False, "error": str(exc)}, status=400)
    except Exception:
        logger.exception("Material import failed")
        messages.error(request, "Ошибка обработки файла материалов")
        if wants_json:
            return JsonResponse({"ok": False, "error": "Ошибка обработки файла материалов"}, status=400)
    return redirect("records", key="materials")


@login_required
@require_POST
def reserve_material(request, pk):
    access(request, "requirements", write=True)
    try:
        quantity = Decimal(request.POST.get("quantity", "0").replace(",", "."))
    except Exception:
        messages.error(request, "Укажите количество резерва")
        return redirect("records", key="requirements")
    if not quantity.is_finite() or quantity == 0:
        messages.error(request, "Количество не может быть нулевым")
        return redirect("records", key="requirements")
    with transaction.atomic():
        requirement = get_object_or_404(models.MaterialRequirement.objects.select_for_update().select_related("order_item"), pk=pk)
        material = models.Material.objects.select_for_update().get(pk=requirement.material_id)
        new_reserved = requirement.reserved + quantity
        if new_reserved < 0 or new_reserved > requirement.required:
            messages.error(request, "Резерв должен быть от нуля до потребности заказа")
            return redirect("records", key="requirements")
        if quantity > material.available_quantity:
            messages.error(request, "Недостаточно доступного материала")
            return redirect("records", key="requirements")
        before = snapshot(requirement)
        requirement.reserved = new_reserved
        material.reserved_quantity += quantity
        requirement.save(update_fields=["reserved"])
        material.save(update_fields=["reserved_quantity"])
        models.AuditLog.objects.create(user=request.user, model_name="MaterialRequirement", object_id=pk,
            action="reserve", reason=request.POST.get("reason", "Резерв под заказ"), before=before, after=snapshot(requirement))
    messages.success(request, "Резерв обновлен")
    return redirect("records", key="requirements")


def pdf_response(title, headers, rows, filename):
    output = io.BytesIO()
    font = cyrillic_pdf_font()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CyrTitle", fontName=font, fontSize=15, leading=19, spaceAfter=15))
    styles.add(ParagraphStyle(name="CyrCell", fontName=font, fontSize=8, leading=11))
    p = lambda v: Paragraph(escape(str(v)), styles["CyrCell"])
    data = [[p(v) for v in headers]] + [[p(v) for v in row] for row in rows]
    widths = [(landscape(A4)[0] - 60) / len(headers)] * len(headers)
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dcebf5")),
        ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#cad8e3")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    doc.build([Paragraph(escape(title), styles["CyrTitle"]), table])
    output.seek(0)
    return FileResponse(output, as_attachment=True, filename=filename, content_type="application/pdf")


@login_required
def print_job(request, pk):
    access(request, "jobs")
    job = get_object_or_404(models.Job.objects.select_related("operation", "employee", "program_line__order_item__order", "program_line__order_item__product"), pk=pk)
    item = job.program_line.order_item
    rows = [
        ("Заказ", item.order.number), ("Изделие", item.product.name),
        ("Партия", item.batch), ("Операция", job.operation.name),
        ("Рабочий центр", job.operation.work_center.name), ("Исполнитель", job.employee or "Не назначен"),
        ("Количество", job.quantity), ("Норма на партию, ч", job.planned_hours),
        ("Штучно-калькуляционное время, ч", round(job.operation.calculated_piece_hours(job.quantity), 3)),
        ("Плановый срок", job.planned_end), ("Состояние", job.get_status_display()),
    ]
    return pdf_response(f"Производственное задание ПЗ-{job.pk}", ["Поле", "Значение"], rows, f"job-{job.pk}.pdf")


@login_required
def print_route(request, pk):
    access(request, "routes")
    product = get_object_or_404(models.Product, pk=pk)
    item = None
    if request.GET.get("item"):
        item = get_object_or_404(models.OrderItem, pk=request.GET["item"], product=product)
    card = route_card_data(product, item)
    return FileResponse(render_route_card(card), as_attachment=True,
        filename=f"route-gost-3.1118-{pk}.pdf", content_type="application/pdf")


@login_required
def print_operation_cards(request, pk):
    access(request, "routes")
    product = get_object_or_404(models.Product, pk=pk)
    rows = []
    for op in product.route.select_related("work_center", "equipment", "definition").order_by("number"):
        norm = getattr(op, "cutting_norm", None)
        rows.append((op.number, op.name, op.get_kind_display(),
            op.equipment.name if op.equipment else op.work_center.name,
            op.setup_hours, op.effective_piece_hours,
            f"n={norm.spindle_rpm:.0f} об/мин; F={norm.feed_mm_min:.1f} мм/мин" if norm and norm.approved else "Режим не утвержден"))
    return pdf_response(f"Проект операционных карт · {product.name} · {product.drawing}",
        ["№", "Операция", "Вид", "Оборудование", "Тпз, ч", "Тшт, ч", "Режим"], rows,
        f"operation-cards-{pk}.pdf")


@login_required
def print_tooling_sheet(request, pk):
    access(request, "routes")
    product = get_object_or_404(models.Product, pk=pk)
    rows = []
    for op in product.route.select_related("equipment").order_by("number"):
        norm = getattr(op, "cutting_norm", None)
        rows.append((op.number, op.name, op.equipment.name if op.equipment else "Не назначено",
            norm.tool if norm else "Не задан", op.machine_setup or "Не задана", op.cnc_program or "—"))
    return pdf_response(f"Проект ведомости оснастки · {product.name} · {product.drawing}",
        ["№", "Операция", "Оборудование", "Инструмент", "Наладка", "Программа"], rows,
        f"tooling-sheet-{pk}.pdf")


@login_required
def print_kit_sheet(request, pk):
    access(request, "requirements")
    item = get_object_or_404(models.OrderItem.objects.select_related("order", "product"), pk=pk)
    rows = [(requirement.material.code, requirement.material.name,
             requirement.quantity_per_item, requirement.required, requirement.material.unit)
            for requirement in models.MaterialRequirement.objects.filter(order_item=item).select_related("material")]
    return pdf_response(f"Проект комплектовочной ведомости · {item.order.number} · {item.product.name}",
        ["Код", "Материал", "На изделие", "На партию", "Ед."], rows,
        f"kit-sheet-{pk}.pdf")


@login_required
def export_report_pdf(request):
    if role(request.user) not in REPORT_ROLES and not request.user.is_superuser:
        raise PermissionDenied
    rows = [(o.number, o.customer.name, o.department.name, o.quantity, o.accepted_quantity,
             f"{o.progress}%", o.due_date, o.forecast_date, o.get_status_display()) for o in report_queryset(request).select_related("customer", "department")]
    return pdf_response("Портфель заказов", ["Заказ", "Заказчик", "Подразделение", "Объем", "Принято", "Готовность", "Срок", "Прогноз", "Статус"], rows, "portfolio.pdf")


@login_required
def reports(request):
    if not request.user.is_superuser and role(request.user) not in REPORT_ROLES:
        raise PermissionDenied
    orders = report_queryset(request)
    rows = []
    for order in orders.select_related("customer", "department"):
        planned_hours = sum((line.planned_hours for line in models.ProgramLine.objects.filter(order_item__order=order, program__status__in=["approved", "revised"]).select_related("order_item__product")), Decimal("0"))
        actual_hours = models.Job.objects.filter(program_line__order_item__order=order).aggregate(v=Sum("actual_hours"))["v"] or Decimal("0")
        rejected = order.events.filter(kind="rejected").aggregate(v=Sum("quantity"))["v"] or 0
        rows.append((order, planned_hours, actual_hours, rejected))
    return render(request, "production/reports.html", context(request, rows=rows,
        departments=models.Department.objects.filter(archived=False)))


def report_queryset(request):
    orders = models.Order.objects.filter(archived=False)
    if request.GET.get("department"):
        orders = orders.filter(department_id=request.GET["department"])
    if request.GET.get("from"):
        orders = orders.filter(due_date__gte=request.GET["from"])
    if request.GET.get("to"):
        orders = orders.filter(due_date__lte=request.GET["to"])
    return orders


@login_required
def export_report(request):
    if role(request.user) not in REPORT_ROLES and not request.user.is_superuser:
        raise PermissionDenied
    wb = Workbook()
    ws = wb.active
    ws.title = "Портфель заказов"
    ws.append(["Заказ", "Заказчик", "Подразделение", "Количество", "Принято", "Готовность %", "Срок", "Прогноз", "Статус"])
    for order in report_queryset(request).select_related("customer", "department"):
        ws.append([spreadsheet_value(value) for value in [order.number, order.customer.name, order.department.name, order.quantity,
            order.accepted_quantity, order.progress, order.due_date.isoformat(), order.forecast_date.isoformat(), order.get_status_display()]])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0865A8")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in "ABCDEFGHI":
        ws.column_dimensions[col].width = 22
    data = io.BytesIO()
    wb.save(data)
    data.seek(0)
    models.AuditLog.objects.create(user=request.user, model_name="Report", object_id=0, action="export")
    return FileResponse(data, as_attachment=True, filename="portfolio.xlsx", content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@login_required
def users(request):
    if role(request.user) != "admin" and not request.user.is_superuser:
        raise PermissionDenied
    return render(request, "production/users.html", context(request, users=User.objects.prefetch_related("groups").order_by("username")))


@login_required
def user_form(request, pk=None):
    if role(request.user) != "admin" and not request.user.is_superuser:
        raise PermissionDenied
    instance = get_object_or_404(User, pk=pk) if pk else None
    form = UserForm(request.POST or None, instance=instance)
    if request.method == "POST" and form.is_valid():
        if not instance and not form.cleaned_data["password"]:
            form.add_error("password", "Для нового пользователя нужен пароль")
        else:
            user = form.save()
            models.AuditLog.objects.create(user=request.user, model_name="User", object_id=user.pk,
                action="update" if instance else "create", reason=request.POST.get("change_reason", ""))
            return redirect("users")
    return render(request, "production/form_page.html", context(request, form=form, title="Пользователь", key="users"))


@login_required
def audit(request):
    if role(request.user) not in ("admin", "manager") and not request.user.is_superuser:
        raise PermissionDenied
    return render(request, "production/audit.html", context(request, entries=models.AuditLog.objects.select_related("user")[:300]))
